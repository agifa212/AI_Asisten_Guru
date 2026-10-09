import os
import io
import streamlit as st
from openai import OpenAI
from docx import Document


# =====================================
# KONFIGURASI
# =====================================

st.set_page_config(
    page_title="AI Asisten Guru By Agifa",
    page_icon="🤖",
    layout="centered"
)


# =====================================
# API KEY
# =====================================

API_KEY = os.getenv("OPENROUTER_API_KEY")

if not API_KEY:
    st.error("❌ API Key OpenRouter belum ditemukan.")
    st.stop()


client = OpenAI(
    base_url="https://openrouter.ai/api/v1",
    api_key=API_KEY
)


# =====================================
# HEADER
# =====================================

st.title("🤖 AI Asisten Guru By Agifa")

st.write(
    "Membantu guru membuat materi, soal, "
    "LKPD, dan perangkat pembelajaran dengan AI."
)

st.divider()


# =====================================
# DATA GURU
# =====================================

nama_guru = st.text_input(
    "👨‍🏫 Nama Guru",
    placeholder="Contoh: Achmad Ginanjar, S.E."
)


# =====================================
# JENJANG
# =====================================

jenjang = st.selectbox(
    "🏫 Jenjang Pendidikan",
    [
        "SD / MI",
        "SMP / MTs",
        "SMA / MA"
    ]
)


# =====================================
# KELAS
# =====================================

if jenjang == "SD / MI":

    kelas = st.selectbox(
        "🎓 Kelas",
        ["1", "2", "3", "4", "5", "6"]
    )

elif jenjang == "SMP / MTs":

    kelas = st.selectbox(
        "🎓 Kelas",
        ["7", "8", "9"]
    )

else:

    kelas = st.selectbox(
        "🎓 Kelas",
        ["10", "11", "12"]
    )


# =====================================
# MATA PELAJARAN
# =====================================

mata_pelajaran = st.text_input(
    "📚 Mata Pelajaran",
    placeholder="Contoh: Ekonomi / Nahwu / Sharaf / Fiqih"
)


# =====================================
# JENIS OUTPUT
# =====================================

jenis_tugas = st.selectbox(
    "🤖 Apa yang ingin dibuat?",
    [
        "📚 Materi Pembelajaran",
        "📝 Bank Soal",
        "📋 Kisi-kisi",
        "🎯 Tujuan Pembelajaran",
        "💡 Pertanyaan Pemantik",
        "🎮 Aktivitas Pembelajaran",
        "📄 LKPD",
        "📑 RPP",
        "📘 Silabus",
        "📊 Asesmen & Rubrik"
    ]
)


# =====================================
# PENGATURAN BANK SOAL
# =====================================

if jenis_tugas == "📝 Bank Soal":

    st.subheader("⚙️ Pengaturan Bank Soal")

    jumlah_soal = st.selectbox(
        "🔢 Jumlah Soal",
        [5, 10, 20, 30]
    )

    tingkat_kesulitan = st.selectbox(
        "📊 Tingkat Kesulitan",
        [
            "Mudah",
            "Sedang",
            "Sulit",
            "Campuran"
        ]
    )

    bentuk_soal = st.selectbox(
        "📝 Bentuk Soal",
        [
            "Pilihan Ganda",
            "Essay"
        ]
    )

else:

    jumlah_soal = 10
    tingkat_kesulitan = "Sedang"
    bentuk_soal = "Pilihan Ganda"


# =====================================
# PENGATURAN LKPD
# =====================================

if jenis_tugas == "📄 LKPD":

    st.subheader("📄 Pengaturan LKPD")

    jenis_pembelajaran = st.selectbox(
        "🏫 Jenis Pembelajaran",
        [
            "Sekolah",
            "Pondok/Pesantren"
        ]
    )

    alokasi_waktu = st.text_input(
        "⏱️ Alokasi Waktu",
        value="2 x 45 menit"
    )

    bentuk_aktivitas = st.multiselect(
        "🎯 Bentuk Aktivitas",
        [
            "Membaca",
            "Mengamati",
            "Menjawab Pertanyaan",
            "Diskusi",
            "Latihan",
            "Studi Kasus",
            "Praktik",
            "Refleksi"
        ],
        default=[
            "Membaca",
            "Menjawab Pertanyaan",
            "Latihan",
            "Refleksi"
        ]
    )

else:

    jenis_pembelajaran = "Sekolah"
    alokasi_waktu = "2 x 45 menit"
    bentuk_aktivitas = []


# =====================================
# PENGATURAN RPP
# =====================================

if jenis_tugas == "📑 RPP":

    st.subheader("📑 Pengaturan RPP")

    jenis_pembelajaran_rpp = st.selectbox(
        "🏫 Jenis Pembelajaran",
        [
            "Sekolah",
            "Pondok/Pesantren"
        ]
    )

    alokasi_waktu_rpp = st.text_input(
        "⏱️ Alokasi Waktu RPP",
        value="2 x 45 menit"
    )

    metode_pembelajaran_rpp = st.multiselect(
        "👨‍🏫 Metode Pembelajaran",
        [
            "Ceramah Interaktif",
            "Diskusi",
            "Tanya Jawab",
            "Problem Based Learning",
            "Project Based Learning",
            "Cooperative Learning",
            "Demonstrasi",
            "Praktik",
            "Studi Kasus"
        ],
        default=[]
    )

else:

    jenis_pembelajaran_rpp = "Sekolah"
    alokasi_waktu_rpp = "2 x 45 menit"
    metode_pembelajaran_rpp = []


# =====================================
# PENGATURAN SILABUS
# =====================================

if jenis_tugas == "📘 Silabus":

    st.subheader("📘 Pengaturan Silabus")

    jenis_pembelajaran_silabus = st.selectbox(
        "🏫 Jenis Pembelajaran",
        [
            "Sekolah",
            "Pondok/Pesantren"
        ]
    )

    jumlah_pertemuan = st.number_input(
        "📅 Jumlah Pertemuan",
        min_value=1,
        max_value=40,
        value=4,
        step=1
    )

    alokasi_waktu_silabus = st.text_input(
        "⏱️ Alokasi Waktu per Pertemuan",
        value="2 x 45 menit"
    )

else:

    jenis_pembelajaran_silabus = "Sekolah"
    jumlah_pertemuan = 4
    alokasi_waktu_silabus = "2 x 45 menit"


# =====================================
# PENGATURAN ASESMEN & RUBRIK
# =====================================

if jenis_tugas == "📊 Asesmen & Rubrik":

    st.subheader("📊 Pengaturan Asesmen & Rubrik")

    jenis_asesmen = st.selectbox(
        "📝 Jenis Asesmen",
        [
            "Asesmen Pengetahuan",
            "Asesmen Keterampilan",
            "Asesmen Sikap",
            "Rubrik Penilaian"
        ]
    )

    bentuk_asesmen = st.selectbox(
        "📋 Bentuk Penilaian",
        [
            "Pilihan Ganda",
            "Essay",
            "Tugas",
            "Praktik",
            "Presentasi",
            "Proyek",
            "Produk",
            "Observasi"
        ]
    )

    skala_penilaian = st.selectbox(
        "📊 Skala Penilaian",
        [
            "1–4",
            "1–5",
            "0–100"
        ]
    )

else:

    jenis_asesmen = "Asesmen Pengetahuan"
    bentuk_asesmen = "Tugas"
    skala_penilaian = "1–4"



# =====================================
# MATERI
# =====================================

materi = st.text_area(
    "📖 Materi Pembelajaran",
    placeholder=(
        "Contoh sekolah: Kelangkaan dan kebutuhan manusia\n"
        "Contoh pondok: Fi'il Madhi dalam Ilmu Sharaf"
    )
)


# =====================================
# KARAKTERISTIK SISWA
# =====================================

def karakteristik_siswa(jenjang):

    if jenjang == "SD / MI":

        return """
Siswa sekolah dasar.

Gunakan bahasa sangat sederhana,
konkret, komunikatif, dan dekat dengan
kehidupan sehari-hari anak.

Hindari istilah akademik yang terlalu sulit.
"""

    elif jenjang == "SMP / MTs":

        return """
Siswa sekolah menengah pertama.

Gunakan bahasa sederhana tetapi mulai
menggunakan istilah akademik.

Dorong pemahaman konsep, hubungan
sebab-akibat, dan penerapan dalam
kehidupan sehari-hari.
"""

    else:

        return """
Siswa sekolah menengah atas atau madrasah
aliyah.

Gunakan bahasa akademik yang tetap mudah
dipahami.

Dorong kemampuan berpikir kritis,
analisis, penerapan konsep, dan
pemecahan masalah.
"""


# =====================================
# BUAT PROMPT
# =====================================

def buat_prompt():

    karakter = karakteristik_siswa(jenjang)

    # ---------------------------------
    # BANK SOAL
    # ---------------------------------

    if jenis_tugas == "📝 Bank Soal":

        instruksi = f"""
Buatlah bank soal sebanyak {jumlah_soal} soal.

Bentuk soal:
{bentuk_soal}

Tingkat kesulitan:
{tingkat_kesulitan}

Jika Pilihan Ganda:

- Gunakan pilihan A, B, C, D.
- Hanya satu jawaban benar.
- Buat pengecoh yang masuk akal.
- Sertakan kunci jawaban.
- Sertakan pembahasan singkat.

Jika Essay:

- Buat pertanyaan yang jelas.
- Sertakan kunci/poin jawaban.
- Sertakan pedoman penilaian.
"""

    # ---------------------------------
    # LKPD
    # ---------------------------------

    elif jenis_tugas == "📄 LKPD":

        konteks = (
            "pembelajaran sekolah"
            if jenis_pembelajaran == "Sekolah"
            else "pembelajaran pondok/pesantren"
        )

        instruksi = f"""
Buatkan LKPD untuk {konteks}.

Alokasi waktu:
{alokasi_waktu}

Bentuk aktivitas:
{", ".join(bentuk_aktivitas)}

Struktur LKPD:

1. Judul LKPD
2. Identitas pembelajaran
3. Tujuan pembelajaran
4. Petunjuk penggunaan LKPD
5. Ringkasan materi
6. Aktivitas peserta didik
7. Latihan atau pertanyaan
8. Refleksi
9. Kesimpulan
10. Kunci jawaban jika terdapat soal yang memiliki jawaban pasti

Sesuaikan tingkat bahasa dan tingkat kesulitan
dengan jenjang dan kelas siswa.

Jika pembelajaran merupakan pondok/pesantren:

- Gunakan istilah yang sesuai dengan materi pesantren.
- Sesuaikan aktivitas dengan karakter pembelajaran
  keislaman atau kepesantrenan.
- Jangan memaksakan format kurikulum sekolah
  apabila materi merupakan materi khas pesantren.
- Tetap buat aktivitas yang praktis dan dapat
  dikerjakan oleh santri.

LKPD harus praktis dan benar-benar dapat
digunakan guru di kelas.
"""

    # ---------------------------------
    # RPP
    # ---------------------------------

    elif jenis_tugas == "📑 RPP":

        konteks_rpp = (
            "pembelajaran sekolah"
            if jenis_pembelajaran_rpp == "Sekolah"
            else "pembelajaran pondok/pesantren"
        )

        instruksi = f"""
Buatkan RPP untuk {konteks_rpp}.

Alokasi waktu:
{alokasi_waktu_rpp}

Metode pembelajaran:
{", ".join(metode_pembelajaran_rpp)}

Susun RPP secara sistematis dengan struktur:

1. Identitas Pembelajaran
2. Tujuan Pembelajaran
3. Materi Pembelajaran
4. Metode Pembelajaran
5. Media dan Sumber Belajar
6. Langkah-Langkah Pembelajaran

A. Kegiatan Pendahuluan
- Salam dan pembukaan
- Apersepsi
- Motivasi
- Penyampaian tujuan pembelajaran

B. Kegiatan Inti
- Aktivitas guru
- Aktivitas peserta didik
- Kegiatan memahami materi
- Kegiatan latihan atau penerapan
- Diskusi atau pemecahan masalah

C. Kegiatan Penutup
- Kesimpulan
- Refleksi
- Evaluasi
- Tindak lanjut

7. Asesmen/Penilaian
- Penilaian sikap
- Penilaian pengetahuan
- Penilaian keterampilan

8. Remedial
9. Pengayaan
10. Refleksi Guru

Sesuaikan RPP dengan jenjang, kelas,
mata pelajaran, materi, dan karakteristik siswa.

Jika merupakan pembelajaran pondok/pesantren,
gunakan pendekatan yang sesuai dengan karakter
pendidikan pesantren dan jangan memaksakan
format sekolah secara kaku.

RPP harus praktis dan dapat digunakan guru
sebagai panduan mengajar di kelas.
"""


    # ---------------------------------
    # SILABUS
    # ---------------------------------

    elif jenis_tugas == "📘 Silabus":

        konteks_silabus = (
            "pembelajaran sekolah"
            if jenis_pembelajaran_silabus == "Sekolah"
            else "pembelajaran pondok/pesantren"
        )

        instruksi = f"""
Buatkan silabus untuk {konteks_silabus}.

Jumlah pertemuan:
{jumlah_pertemuan}

Alokasi waktu setiap pertemuan:
{alokasi_waktu_silabus}

Susun silabus secara sistematis dan praktis.

Struktur silabus:

1. Identitas Pembelajaran
   - Nama sekolah/pesantren
   - Mata pelajaran
   - Jenjang
   - Kelas
   - Semester

2. Tujuan Pembelajaran

3. Materi Pembelajaran

4. Rencana Pembelajaran per Pertemuan

Untuk setiap pertemuan tampilkan:
- Pertemuan ke-
- Materi
- Tujuan/kompetensi yang ingin dicapai
- Kegiatan pembelajaran
- Metode
- Alokasi waktu
- Bentuk asesmen
- Sumber belajar

5. Media Pembelajaran

6. Sumber Belajar

7. Bentuk Asesmen

8. Remedial

9. Pengayaan

Buat silabus yang realistis dan dapat digunakan
guru sebagai acuan pembelajaran.

Sesuaikan dengan jenjang, kelas, mata pelajaran,
materi, dan karakteristik siswa.

Jika merupakan pembelajaran pondok/pesantren,
gunakan pendekatan yang sesuai dengan karakter
pendidikan pesantren.

Jangan memaksakan istilah atau struktur kurikulum
sekolah apabila materi merupakan materi khas
pondok/pesantren.
"""


    # ---------------------------------
    # ASESMEN & RUBRIK
    # ---------------------------------

    elif jenis_tugas == "📊 Asesmen & Rubrik":

        instruksi = f"""
Buatkan {jenis_asesmen} untuk pembelajaran
{mata_pelajaran} kelas {kelas}.

Bentuk penilaian:
{bentuk_asesmen}

Skala penilaian:
{skala_penilaian}

Materi:
{materi}

Buat instrumen penilaian yang praktis,
jelas, objektif, dan dapat langsung digunakan
oleh guru.

Jika Asesmen Pengetahuan:

Buat instrumen untuk mengukur pemahaman siswa
terhadap materi.

Jika Asesmen Keterampilan:

Buat tugas/praktik/proyek yang dapat mengukur
kemampuan siswa menerapkan materi.

Jika Asesmen Sikap:

Buat lembar observasi dengan aspek seperti:
- Disiplin
- Tanggung jawab
- Kerja sama
- Keaktifan
- Sikap selama pembelajaran

Jika Rubrik Penilaian:

Buat tabel rubrik dengan struktur:

No | Aspek yang Dinilai | Bobot | Skor 4 | Skor 3 | Skor 2 | Skor 1

Sesuaikan aspek penilaian dengan tugas
yang diberikan.

Jelaskan kriteria setiap tingkat skor
secara jelas.

Jika menggunakan bobot, pastikan total
bobot = 100%.

Tambahkan:
1. Petunjuk penggunaan
2. Instrumen penilaian
3. Kriteria penilaian
4. Rumus/perhitungan nilai
5. Interpretasi hasil

Jika pembelajaran merupakan pondok/pesantren,
sesuaikan aspek penilaian dengan karakter
materi pesantren.

Contoh:

Untuk Ilmu Sharaf:
- Ketepatan tashrif
- Pemahaman kaidah
- Ketepatan penerapan

Untuk Tahfizh:
- Kelancaran
- Ketepatan hafalan
- Tajwid
- Makhraj

Jangan memaksakan indikator sekolah
pada materi khas pesantren.
"""


    # ---------------------------------
    # OUTPUT LAIN
    # ---------------------------------

    else:

        instruksi = f"""
Buatlah {jenis_tugas} berdasarkan materi.

Sesuaikan dengan:
- Jenjang pendidikan
- Kelas
- Mata pelajaran
- Karakteristik siswa

Gunakan struktur yang sistematis dan mudah
digunakan oleh guru.
"""

    # ---------------------------------
    # PROMPT UTAMA
    # ---------------------------------

    return f"""
Anda adalah AI Asisten Guru profesional.

Nama Guru:
{nama_guru}

Jenjang:
{jenjang}

Kelas:
{kelas}

Mata Pelajaran:
{mata_pelajaran}

Jenis Output:
{jenis_tugas}

Jenis Pembelajaran:
{jenis_pembelajaran}

Materi:
{materi}

Karakteristik siswa:

{karakter}

Instruksi:

{instruksi}

Gunakan bahasa Indonesia yang jelas,
edukatif, sistematis, dan sesuai tingkat
pemahaman siswa.

Jangan keluar dari materi yang diberikan.
"""


# =====================================
# PANGGIL AI
# =====================================

def panggil_ai(prompt):

    response = client.chat.completions.create(

        model="openrouter/free",

        messages=[

            {
                "role": "system",
                "content": """
Anda adalah AI Asisten Guru profesional.

Bantu guru membuat perangkat pembelajaran
yang akurat, sistematis, praktis, dan sesuai
dengan jenjang pendidikan siswa.

Jika materi berasal dari pondok/pesantren,
hormati karakteristik pembelajaran pesantren.
"""
            },

            {
                "role": "user",
                "content": prompt
            }

        ]
    )

    return response.choices[0].message.content


# =====================================
# BUAT FILE WORD
# =====================================

def buat_word(judul, isi):

    document = Document()

    document.add_heading(
        judul,
        level=1
    )

    document.add_paragraph(
        f"Nama Guru: {nama_guru}"
    )

    document.add_paragraph(
        f"Jenjang: {jenjang}"
    )

    document.add_paragraph(
        f"Kelas: {kelas}"
    )

    document.add_paragraph(
        f"Mata Pelajaran: {mata_pelajaran}"
    )

    document.add_paragraph("")

    document.add_heading(
        "Hasil Pembelajaran",
        level=2
    )

    for baris in isi.split("\n"):

        baris = baris.strip()

        if baris:

            document.add_paragraph(baris)

    file = io.BytesIO()

    document.save(file)

    file.seek(0)

    return file


# =====================================
# GENERATE
# =====================================

if st.button(
    "🚀 Generate",
    use_container_width=True
):

    if not nama_guru:

        st.warning(
            "⚠️ Nama guru belum diisi."
        )

    elif not mata_pelajaran:

        st.warning(
            "⚠️ Mata pelajaran belum diisi."
        )

    elif not materi:

        st.warning(
            "⚠️ Materi pembelajaran belum diisi."
        )

    else:

        prompt = buat_prompt()

        with st.spinner(
            "🤖 AI sedang menyusun..."
        ):

            try:

                hasil = panggil_ai(prompt)

                st.session_state["hasil_ai"] = hasil
                st.session_state["hasil_terakhir_disimpan"] = False

                st.success(
                    "✅ Hasil berhasil dibuat!"
                )

            except Exception as e:

                st.error(
                    "❌ Terjadi kesalahan saat "
                    "menghubungi AI."
                )

                st.code(str(e))


# =====================================
# TAMPILKAN HASIL
# =====================================

if st.session_state.get("hasil_ai"):

    st.divider()

    st.subheader(
        "📚 Hasil AI Asisten Guru"
    )

    st.markdown(
        st.session_state.get("hasil_ai", "")
    )

    # =====================================
    # SIMPAN KE RIWAYAT
    # =====================================

    if "riwayat" not in st.session_state:
        st.session_state["riwayat"] = []

    if not st.session_state.get(
        "hasil_terakhir_disimpan",
        False
    ):

        st.session_state["riwayat"].append({
            "jenis": jenis_tugas,
            "mata_pelajaran": mata_pelajaran,
            "kelas": kelas,
            "materi": materi,
            "hasil": st.session_state.get(
                "hasil_ai",
                ""
            )
        })

        st.session_state["hasil_terakhir_disimpan"] = True

    # =====================================
    # DOWNLOAD WORD
    # =====================================

    file_word = buat_word(
        f"{jenis_tugas} - {mata_pelajaran} Kelas {kelas}",
        st.session_state.get("hasil_ai", "")
    )

    st.download_button(
        label="📥 Download Word",
        data=file_word,
        file_name=(
            f"AI_Asisten_Guru_"
            f"{mata_pelajaran}_"
            f"Kelas_{kelas}.docx"
        ),
        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        ),
        use_container_width=True
    )


# =====================================
# RIWAYAT DOKUMEN
# =====================================

st.divider()

st.subheader("📚 Riwayat Dokumen")

if st.session_state.get("riwayat"):

    for dokumen in reversed(
        st.session_state["riwayat"]
    ):

        with st.expander(
            f"{dokumen['jenis']} — "
            f"{dokumen['mata_pelajaran']} "
            f"Kelas {dokumen['kelas']}"
        ):

            st.write(
                f"**Materi:** {dokumen['materi']}"
            )

            st.markdown(
                dokumen["hasil"]
            )

else:

    st.info(
        "Belum ada dokumen yang dibuat."
    )